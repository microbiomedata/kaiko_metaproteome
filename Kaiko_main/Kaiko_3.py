## This is originally ExtractTopRankTaxaFromUniRef.py

import pandas as pd
import numpy as np
import time
import re

from pathlib import Path, PureWindowsPath
from openpyxl import load_workbook


# @profile    
def run_diamond_tally(diamond_output, filterby, member_csv, mode, fout, detailed_fout, DB):
    # taxa_stats = pd.read_csv(taxa_stats_path, sep = '\t')

    if mode=="member":
            taxa_col = 'member_taxa'
    elif mode=="common":
        taxa_col = 'common_taxa'
    else:
        print("  [ERR] please select the mode to be either 'member' or 'common'.")
        return
    
    if not detailed_fout.exists():
                
        ############################################################
        # read diamond output file
        ############################################################
        print("Reading diamond output file...")
        dmd = read_dmd(diamond_output)

        ############################################################
        # filter by quality and taxa
        ############################################################
        print("Filtering by quality and taxa...")
        filtered_dmd = dmd_filter(dmd, filterby={})
        if DB == 'Uniref100':
            db_pattern = 'TaxID'
        elif DB == 'reference_proteomes':
            db_pattern = 'OX'
        filtered_dmd = collect_taxid(filtered_dmd, db_pattern)
        filtered_dmd['uniref_id'] = [i[0] for i in filtered_dmd.uniref_seq.str.split(" ", n = 1)]

        ############################################################
        # garbage collection
        ############################################################
        print("Garbage collection...")
        del dmd
        import gc
        gc.collect()
        
        ############################################################
        # retrieve UniRef100 representative taxa and its members
        ############################################################
        # print("Retrieving taxa lineage information...")
        if db_pattern == 'TaxID':
            # unique_unirefs = set(filtered_dmd.uniref_id.drop_duplicates())
            unique_unirefs = {x:1 for x in filtered_dmd.uniref_id.drop_duplicates().tolist()}
            taxa_member_tbl = get_taxa_members(member_csv, unique_unirefs)  
            ############################################################
            # merge
            ############################################################
            print("Adding taxa members...")
            merged = filtered_dmd.merge(taxa_member_tbl, left_on="uniref_id", right_on="uid", how="inner")
            print("  Final table size:{}".format(merged.shape))
            
            if merged.shape[0]!=filtered_dmd.shape[0]:
                print("[WARN] You might use the different version of UniRef100.fasta and .xml")
            
            # get the member taxa of each (scan, uid)
            unique_members = []
            scanids = merged.scans.tolist()
            uids = merged.uid.tolist()
            commons = merged.common_taxa.tolist()
            pidents = merged.pident.tolist()
            align_lens = merged.align_len.tolist()
            scores = merged.score.tolist()
            for ii, members in enumerate(merged.members.str.split(":").tolist()):
                for mm in members:
                    unique_members.append({"scan":scanids[ii], "uid":uids[ii], "protein":uids[ii].replace('UniRef100_', ''),"member_taxa":int(mm), "common_taxa":int(commons[ii]), 
                                           "pident":pidents[ii], "align_len":align_lens[ii], "score":scores[ii]})
            print("  #members:{}".format(len(unique_members)))
            members_df = pd.DataFrame(unique_members)

        elif db_pattern == 'OX':
            members_df = filtered_dmd[['scans']].copy()
            members_df['scan'] = members_df['scans']
            members_df['uid'] = [uniref_id.split('|')[1] for uniref_id in filtered_dmd['uniref_id']]
            members_df['protein'] = members_df['uid']
            members_df['member_taxa'] = filtered_dmd['taxid']
            members_df['common_taxa'] = members_df['member_taxa']
            members_df['pident'] = filtered_dmd['pident']
            members_df['align_len'] = filtered_dmd['align_len']
            members_df['score'] = filtered_dmd['score']
            members_df = pd.DataFrame(members_df)
 
        ############################################################
        # top-rank taxa
        ############################################################
        
        detailed_output = members_df
        # detailed_output["protein"] = [detailed_output.uid[i].replace('UniRef100_', '') for i in range(0, len(detailed_output))]
        detailed_output = detailed_output[['scan', 'uid', 'protein', 'member_taxa', 'common_taxa', 'pident', 'align_len', 'score']]
        # detailed_output = detailed_output.merge(taxa_stats, left_on = taxa_col, right_on = 'taxid', how = 'left')
        detailed_output.to_csv(detailed_fout, index = False)
    else:
        print("Loading |scan|protein|taxa| table " + detailed_fout.name + "\n")
        detailed_output = pd.read_csv(detailed_fout)
    
    if fout.exists():
        book = load_workbook(fout)
        writer = pd.ExcelWriter(fout, engine='openpyxl')
        writer.book = book
        sheet_names = book.sheetnames
    else:
        writer = pd.ExcelWriter(fout, engine='xlsxwriter')
        sheet_names = []

    # pidents.sort(reverse=True)
    if 'score' in filterby.keys():
        score = filterby['score']
        sheet_name = f'alignment >= {score} percent'
        detailed_output_sheet = detailed_output[detailed_output['score'] >= score]
    else:
        pident = filterby['pident']
        sheet_name = f'pident >= {pident} percent'
        detailed_output_sheet = detailed_output[detailed_output['pident'] >= pident]
    if sheet_name not in sheet_names:
        df = make_top_taxa_df(detailed_output_sheet, taxa_col)
        df.to_excel(writer, sheet_name = sheet_name)
    writer.close()

def find_smaller_taxa(df, pepcount_taxid, _n_strain_select, index):
    tax_ranks = ['species', 'genus', 'family', 'order', 'class', 'phylum', 'kingdom', 'superkingdom']
    tax_index = 0
    subcount = []
    while len(subcount) == 0 and tax_index < 8:
        tax_rank = tax_ranks[tax_index]
        if tax_rank == 'species':
            rank_value_col = 'tax_name'
        else:
            rank_value_col = tax_rank
        if _n_strain_select == -1:
            subcount = pepcount_taxid[pepcount_taxid[tax_rank] == df[rank_value_col][index]].copy()
        else:
            subcount = pepcount_taxid[pepcount_taxid[tax_rank] == df[rank_value_col][index]].copy().iloc[0:_n_strain_select]
        subcount['running_coverage'] = df['running_coverage'][index]
        subcount['notes'] = 'Secondary taxa identified by Kaiko, with the same {} ({}) as the primary taxa ({}) below the proteome size cutoff. \'hits\' denotes tally1 hits'.format(tax_rank, str(df[rank_value_col][index]), df['tax_name'][index])
        tax_index = tax_index + 1

    return(subcount)

def read_dmd(diamond_output):
    dmd_colnames = ['scans','uniref_seq','pident','evalue','mismatch', 'gapopen', 
                    'gaps', 'qstart', 'qend', 'qseq', 'sstart', 'send', 'sseq', 'qlen', 'full_qseq']
    dmd = pd.read_csv(diamond_output, sep='\t', header=None, names=dmd_colnames)
    dmd['align_len'] = [len(seq) for seq in dmd['qseq']]
    dmd['score'] = [align_len * pident/qlen for align_len, pident, qlen in zip(dmd['align_len'], dmd['pident'], dmd['qlen'])]
    return dmd

def dmd_filter(dmd, filterby={}):
    if len(filterby) > 0:
        filterby_cond = None
        for col in filterby:
            if col in ['pident', 'align_len', 'score']:
                tmp = dmd[col]>=filterby[col]
            elif col in ['evalue', 'mismatch']:
                tmp = dmd[col]<=filterby[col]
            else:
                print("  [ERR] no column found. For filtering, please use 'pident','evalue','mismatch'.")
                raise Exception

            if filterby_cond is None:
                filterby_cond = tmp
            else:
                filterby_cond &= tmp

        filtered_dmd = dmd[filterby_cond].copy()
        print("  org:{}, filtered:{}".format(dmd.shape[0], filtered_dmd.shape[0]))
        return filtered_dmd
    else:
        return dmd

def collect_taxid(filtered_dmd, db_pattern = 'TaxID'):
    filtered_dmd['taxid'] = filtered_dmd.uniref_seq.str.extract(f'^.+? {db_pattern}=(\d*)\s?')
    # filtered_dmd['taxid'] = filtered_dmd.uniref_seq.str.extract('^.+? OX=(\d*)\s?')

    # drop the irrelevant
    # unknown
    # 1  # root
    # 9606  # Homo sapiens
    # 412755  # marine sediment metagenome
    # 408172  # marine metagenome
    # 9823 # Sus scrofa (pig)
    drop_taxids = set(['1','2','9606','412755','408172', '9823'])
    print("Dropping the following taxids\n")
    print(drop_taxids)

    _filtered_dmd = filtered_dmd[~((filtered_dmd.taxid=='')|(filtered_dmd.taxid.isin(drop_taxids)))].copy()
    _filtered_dmd.taxid = _filtered_dmd.taxid.astype(int)
    _filtered_dmd = _filtered_dmd.reset_index(drop=True)
    print('filter relevant taxa:{}'.format(_filtered_dmd.shape[0]))

    wrong_taxids = [444888, 55087, 210425]
    print("{} wrong taxids found in the results and corrected.".format(
        _filtered_dmd[_filtered_dmd.taxid.isin(wrong_taxids)].shape[0]))

    _filtered_dmd.loc[_filtered_dmd.taxid==444888, 'taxid'] = 629  # Yersinia
    _filtered_dmd.loc[_filtered_dmd.taxid==55087, 'taxid'] = 1386  # Bacillus
    _filtered_dmd.loc[_filtered_dmd.taxid==210425, 'taxid'] = 583  # Proteus
    return _filtered_dmd

def get_taxa_members(member_tbl_file, unique_unirefs):
    chunksize = 1000000

    stime = time.time()
    dfs = []
    num_iters = 0
    for chunk in pd.read_csv(member_tbl_file, chunksize=chunksize):
        # if num_iters > 10:
        #     break
        # else:
        tdf = chunk[chunk.uid.isin(unique_unirefs.keys())]
        dfs.append(tdf)
        num_iters += 1

    df = pd.concat(dfs)
    print("  #Chunk:{}, Size:{}, {:.2f}min".format(len(dfs), df.shape, (time.time()-stime)/60))
    return df

def make_top_taxa_df(detailed_output, taxa_col):
    unique_pepseq_taxa = detailed_output.drop_duplicates(subset=['scan',taxa_col])
    pepcount_taxid = unique_pepseq_taxa[taxa_col].value_counts()

    print('  unique peptides: {}'.format(detailed_output.scan.value_counts().shape[0]))
    print('  unique taxa: {}'.format(pepcount_taxid.shape[0]))

    def besthit(row):
        return pepcount_taxid[row[taxa_col]].nlargest(1, keep='all').index.tolist()

    besthits = []
    for besthit in unique_pepseq_taxa.groupby(by='scan').apply(besthit):
        besthits += besthit
    besthits = pd.Series(besthits).value_counts()

    pepcount_taxid = pepcount_taxid.to_frame('hits')
    pepcount_taxid['taxid'] = pepcount_taxid.index
    # pepcount_taxid = pepcount_taxid.merge(taxa_stats, left_on = 'taxid', right_on = 'taxid', how = 'left')
    
    # if n_strain_select > 0 & n_strain_select <= 5:
    #     _n_strain_select = 5
    # elif n_strain_select > 5:
    #     _n_strain_select = n_strain_select
    # else:
    #     _n_strain_select = -1
    
    # top_taxa = besthits.nlargest(n_strain_select, keep='all').to_frame(name='hits')
    top_taxa = besthits.to_frame(name='hits')
    top_taxa['taxid'] = top_taxa.index

    ############################################################
    # save top-rank taxa info
    ############################################################
    print("Saving top-rank taxa info...")
    # df = pd.concat([top_taxa, taxa_stats], join='inner', axis=1)
    # df = top_taxa.merge(taxa_stats, left_on = 'taxid', right_on = 'taxid', how = 'left')
    df = top_taxa
    df['proportion'] = df['hits']/df['hits'].sum()
    df['running_coverage'] = df['hits'].cumsum()/df['hits'].sum()
    append = pd.DataFrame({'taxid' : [],
                        'tax_name' : [],
                        'species' : [],
                        'rank' : [],
                        'hits' : [],
                        'n_protein' : [],
                        'proportion' : [],
                        'running_coverage' : [],
                        'notes' : []})
    df['notes'] = 'Primary taxa identified by Kaiko. \'hits\' denotes tally2 hits'
    # pepcount_taxid = pepcount_taxid[pepcount_taxid['n_protein'] < n_protein_cutoff]
    # for index in range(len(df)):
    #     if df.iloc[index].n_protein > n_protein_cutoff:
    #         subcount = find_smaller_taxa(df, pepcount_taxid, _n_strain_select, index)
    #         append = pd.concat([append, subcount])

    df = pd.concat([df, append])
    df = df.sort_values(by = ['running_coverage', 'notes'])
    return(df)